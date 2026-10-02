import os
import re
import json
import pypdf
from typing import Optional
from openai import OpenAI
from backend.models import InvoiceData

import docx

def extract_raw_text_from_pdf(pdf_path: str) -> str:
    reader = pypdf.PdfReader(pdf_path)
    full_text = []
    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        full_text.append(f"--- PAGE {idx+1} ---\n{text}")
    return "\n\n".join(full_text)

def extract_raw_text_from_docx(docx_path: str) -> str:
    doc = docx.Document(docx_path)
    full_text = []
    for p in doc.paragraphs:
        if p.text.strip():
            full_text.append(p.text.strip())
    # Extract from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
    return "\n".join(full_text)

def extract_raw_text_from_doc(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return extract_raw_text_from_docx(file_path)
    return extract_raw_text_from_pdf(file_path)

SYSTEM_PROMPT = """You are an expert accounting system and invoice parser specialized in European standard EN 16931 and ZUGFeRD / Factur-X specifications.
Your goal is to parse raw text extracted from invoice PDFs into a clean, complete, structured InvoiceData object.

Instructions:
1. Identify Invoice Number, Issue Date (format YYYY-MM-DD or DD.MM.YYYY), Due Date, and Currency (e.g. EUR).
2. Identify Seller (Supplier) details: Name, Street, Postcode, City, Country Code (ISO alpha-2 e.g. DE, SI), VAT ID (e.g. DE814113634), Tax Registration Number (e.g. HRB number or local tax number), and Contact details.
3. Identify Buyer (Customer) details: Name, Street, Postcode, City, Country Code (e.g. SI, DE), VAT ID (e.g. SI52093808), and Buyer Reference (if present).
4. Identify Delivery & Logistics (BG-13): If present, extract destination name, street, postcode, city, country_code, delivery_date (BT-72), location_id (BT-71, e.g. unloading point / Abladestelle number like '960' from 'Abladestelle Nr. 960'), and despatch_advice_id (BT-16).
   STRICT RULE FOR BT-16: Only extract `despatch_advice_id` if an actual delivery note / packing slip number is explicitly present next to 'Lieferschein-Nr.'. If no document number appears after 'Lieferschein-Nr.', return null. NEVER extract words like 'Lieferdatum', 'Lieferschein', or 'Bestell-Nr.' as the reference number. If the field is blank or contains only colons/labels (e.g., service invoices where no delivery note exists), return null (None). NEVER extract label fragments such as '-Nr', 'Nr.', or colons.
5. Payment information: IBAN (remove spaces), BIC/SWIFT, Bank Name, and payment_terms text (e.g. exact terms from 'Zahlungsbedingung:' such as 'innerhalb 30 Tagen ohne Abzug').
6. Line items: For each line item, extract name/description, quantity (default 1.0 if not listed), unit code (e.g. C62 for pieces/units, HUR for hours, or DAY), unit price, net line total amount, tax category code ('S' for standard, 'AE' for reverse charge, 'Z' for zero, 'E' for exempt), and tax rate percent.
7. Tax breakdown: Provide tax basis amount, tax percent, tax amount, and exemption reason if reverse charge (e.g. "Reverse charge / The beneficiary is responsible for the taxation of this amount").
8. Totals: Calculate or extract line_net_total, tax_total, grand_total, and due_payable_amount. Verify that line_net_total + tax_total == grand_total (and due_payable_amount).
9. Confidence scoring: In confidence_scores dictionary, assign 'HIGH', 'MEDIUM', or 'LOW' for each field:
   - 'HIGH' if the field was clearly labeled and explicitly printed on the document.
   - 'MEDIUM' if inferred, calculated from other numbers, or slightly ambiguous.
   - 'LOW' if guessed, defaulted, or hard to distinguish on the invoice.
   Keys: 'invoice_number', 'issue_date', 'currency', 'buyer_reference', 'seller_name', 'seller_vat', 'seller_street', 'seller_city', 'buyer_name', 'buyer_vat', 'buyer_street', 'buyer_city', 'iban', 'bic', 'total_net', 'total_tax', 'total_grand', 'total_due'.
"""

def extract_invoice_data(file_path: str, api_key: Optional[str] = None, system_prompt: Optional[str] = None):
    api_key = api_key or os.environ.get("OPENAI_API_KEY")
    if not api_key:
        raise ValueError("OpenAI API Key is required. Please set OPENAI_API_KEY in environment or pass it in.")
    
    prompt_to_use = (system_prompt and system_prompt.strip()) or SYSTEM_PROMPT

    client = OpenAI(api_key=api_key)
    raw_text = extract_raw_text_from_doc(file_path)
    
    completion = client.beta.chat.completions.parse(
        model="gpt-4o",
        messages=[
            {"role": "system", "content": prompt_to_use},
            {"role": "user", "content": f"Please parse this invoice:\n\n{raw_text}"}
        ],
        response_format=InvoiceData,
    )
    
    parsed_data = completion.choices[0].message.parsed
    usage = completion.usage

    # Deterministic recovery for Delivery Details (BT-71 & BT-16) if raw text contains explicit markers
    if parsed_data:
        # 1. Ship-to location ID / Abladestelle (BT-71)
        m_abla = re.search(r"Abladestelle(?:\s+Nr\.?)?\s*[:\s]*([A-Za-z0-9\-_]+)", raw_text, re.IGNORECASE)
        if m_abla:
            abla_val = m_abla.group(1).strip()
            if not parsed_data.delivery:
                from backend.models import DeliveryDetails
                parsed_data.delivery = DeliveryDetails(location_id=abla_val)
            elif not parsed_data.delivery.location_id:
                parsed_data.delivery.location_id = abla_val

        # 2. Despatch Advice / Lieferschein (BT-16) Sanitizer & Strict Label Filter
        # List of known German/English adjacent labels that must never be accepted as IDs
        banned_labels = [
            "lieferdatum", "lieferschein", "lieferschein-nr", "nr", "-nr", "bestell-nr", 
            "kundennummer", "kontakt", "lieferanten-nr", "datum", "none", "null"
        ]
        
        raw_da = getattr(parsed_data.delivery, "despatch_advice_reference", None) or getattr(parsed_data.delivery, "despatch_advice_id", None)
        if raw_da:
            cleaned = str(raw_da).strip().lstrip("-").lstrip(":").strip()
            # Must not be a banned label, must be at least 3 chars, and must contain at least one digit
            if cleaned.lower() in banned_labels or not any(char.isdigit() for char in cleaned) or len(cleaned) < 3:
                parsed_data.delivery.despatch_advice_id = None
                parsed_data.delivery.despatch_advice_reference = None
            else:
                parsed_data.delivery.despatch_advice_id = cleaned
                parsed_data.delivery.despatch_advice_reference = cleaned
        else:
            # Fallback: only if explicit number follows Lieferschein-Nr. on the SAME line/token
            m_ls = re.search(r"Lieferschein(?:-Nr\.?)?\s*[:]\s*([0-9A-Za-z\-_]+)", raw_text, re.IGNORECASE)
            if m_ls:
                cand = m_ls.group(1).strip().lstrip("-").lstrip(":").strip()
                if cand.lower() not in banned_labels and any(char.isdigit() for char in cand) and len(cand) >= 3:
                    if not parsed_data.delivery:
                        from backend.models import DeliveryDetails
                        parsed_data.delivery = DeliveryDetails(despatch_advice_id=cand, despatch_advice_reference=cand)
                    else:
                        parsed_data.delivery.despatch_advice_id = cand
                        parsed_data.delivery.despatch_advice_reference = cand



        # 3. Clean recipient_name if location_id or Tor was appended to it
        if parsed_data.delivery and parsed_data.delivery.recipient_name and parsed_data.delivery.location_id:
            parsed_data.delivery.recipient_name = re.sub(
                r"\s*Abladestelle(?:\s+Nr\.?)?\s*" + re.escape(parsed_data.delivery.location_id),
                "",
                parsed_data.delivery.recipient_name,
                flags=re.IGNORECASE
            ).strip()

        # 4. Extract genuine performance and billing notices (BT-22: IncludedNote) from invoice body
        if not parsed_data.notes:
            parsed_data.notes = []

        # Filter out any synthetic generator notes
        parsed_data.notes = [n for n in parsed_data.notes if "ZUGFeRD PDF/A-3 Generator" not in n]

        # Scan for known standard German invoice body notices if not already captured
        m_muster = re.search(r"(Wir berechnen\s+100\s*%\s*gem[äa]ß\s+beigef[üu]gter\s+Musterpr[üu]fberichte)", raw_text, re.IGNORECASE)
        if m_muster:
            note_text = m_muster.group(1).strip()
            # Normalize umlauts if needed
            note_text = note_text.replace("gem", "gemäß").replace("Musterprfberichte", "Musterprüfberichte")
            if note_text not in parsed_data.notes:
                parsed_data.notes.append(note_text)

        m_leist = re.search(r"(Rechnungsdatum\s*=\s*Leistungsdatum)", raw_text, re.IGNORECASE)
        if m_leist:
            note_text = m_leist.group(1).strip()
            if note_text not in parsed_data.notes:
                parsed_data.notes.append(note_text)


    
    prompt_tokens = usage.prompt_tokens if usage else 0
    completion_tokens = usage.completion_tokens if usage else 0
    total_tokens = usage.total_tokens if usage else 0
    
    # Official OpenAI gpt-4o pricing: $2.50 / 1M prompt tokens, $10.00 / 1M completion tokens
    cost = (prompt_tokens * (2.50 / 1_000_000)) + (completion_tokens * (10.00 / 1_000_000))
    
    usage_info = {
        "model": "gpt-4o",
        "prompt_tokens": prompt_tokens,
        "completion_tokens": completion_tokens,
        "total_tokens": total_tokens,
        "estimated_cost_usd": round(cost, 6),
        "formatted_cost": f"${cost:.4f}"
    }
    
    return parsed_data, usage_info
